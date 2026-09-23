"""
DOCX Document Loader using python-docx with zip/xml fallback
Extracts text from paragraphs and tables.
"""

import io
import logging
from typing import List, Dict, Any

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Extracts structured text from DOCX bytes.
    Returns a list of dicts: [{"page": 1, "text": "..."}]
    """
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        lines = []

        for para in doc.paragraphs:
            if para.text.strip():
                lines.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    lines.append(" | ".join(row_cells))

        full_text = "\n\n".join(lines)
        logger.info(f"Successfully extracted {len(lines)} blocks from DOCX: {filename}")
        return [{"page": 1, "text": full_text}]
    except ImportError:
        logger.warning("python-docx not installed, using zipfile XML fallback.")
    except Exception as e:
        logger.error(f"Error parsing DOCX with python-docx: {e}")

    # Fallback to standard library zipfile XML reading
    try:
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            xml_content = z.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            # Namespace for wordprocessingML
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            text_nodes = tree.findall(".//w:t", ns)
            text = " ".join([node.text for node in text_nodes if node.text])
            return [{"page": 1, "text": text}]
    except Exception as zip_err:
        logger.error(f"XML Fallback DOCX extraction failed: {zip_err}")
        return [{"page": 1, "text": ""}]
